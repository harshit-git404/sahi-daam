from docx import Document
from datetime import datetime

output_path = r"d:\sahi-daam\Sahi_Daam_Project_Review_Update.docx"

doc = Document()

# Title
p = doc.add_paragraph()
run = p.add_run("Sahi Daam Project Review Update")
run.bold = True
run.font.size = 22

# Metadata
meta = doc.add_paragraph()
meta.add_run(f"Date: {datetime.today().strftime('%d %B %Y')}")
meta.add_run("\nProject Type: Fair-price agriculture and retail transparency application")

doc.add_paragraph("")

# Overview
heading = doc.add_heading("1. Project Overview", level=1)
heading.alignment = 1
p = doc.add_paragraph(
    "Sahi Daam is a consumer-focused fair-price companion application designed to help buyers evaluate fruits and vegetables more intelligently. "
    "The system combines computer-vision freshness assessment, produce classification, and wholesale market data to calculate a more objective and fair retail price. "
    "The goal is to empower consumers with transparency, improve negotiation confidence, and encourage informed purchasing decisions."
)

# Current Progress
heading = doc.add_heading("2. Current Progress and Status", level=1)
for item in [
    "The application architecture is in place with a FastAPI backend and a frontend interface for user interaction.",
    "The backend includes routes for scanning, bargaining support, and sector-wise pricing analysis.",
    "The pricing engine integrates wholesale market references and local data inputs to estimate fair value.",
    "Machine learning modules are included for produce classification and freshness evaluation.",
    "The project has a working API structure and a UI flow that supports product evaluation and price comparison.",
    "The current phase focuses on improving pricing logic and making the output more practical for real-world decision-making."
]:
    doc.add_paragraph(item, style='List Bullet')

# Functional Scope
heading = doc.add_heading("3. Functional Scope", level=1)
for item in [
    "Freshness detection and quality assessment using image and model-based analysis.",
    "Produce classification support for identifying fruits and vegetables.",
    "Sector and market-based pricing insights for reference benchmarking.",
    "Negotiation guidance and fair-price assessment for consumers.",
    "Structured API and frontend integration to provide a smooth user experience."
]:
    doc.add_paragraph(item, style='List Bullet')

# Achievements
heading = doc.add_heading("4. Key Achievements", level=1)
for item in [
    "Established a complete project structure for backend, frontend, ML, and pricing services.",
    "Connected market data and pricing logic to support fair-value estimation.",
    "Developed a clear architecture to keep logic modular and maintainable.",
    "Built the foundation for future product enhancement and real-world deployment.",
    "Created a review-friendly project basis that can be expanded with more advanced pricing parameters."
]:
    doc.add_paragraph(item, style='List Bullet')

# Future Additions
heading = doc.add_heading("5. Planned Future Additions", level=1)
for item in [
    "Distance-wise price calculation: Add a pricing model that considers travel distance, transportation cost, and sourcing location to provide a more realistic market price.",
    "Community price calculation: Add a community-level pricing feature that considers local market sentiment, peer pricing patterns, and neighborhood-based reference values.",
    "Improve price fairness by combining market price, quality score, transport impact, and local community demand data.",
    "Extend analytics to support more accurate and transparent price recommendations for end users."
]:
    doc.add_paragraph(item, style='List Bullet')

# Review Summary
heading = doc.add_heading("6. Review Summary", level=1)
summary = (
    "The project has successfully established a strong base for a fair-price agriculture application with working backend, "
    "frontend, and AI-based modules. Current work demonstrates meaningful progress toward a practical consumer tool. "
    "The next phase should focus on advanced pricing logic, especially distance-wise and community-based calculations, which will make the solution more adaptive, locality-aware, and valuable for real-world use."
)
doc.add_paragraph(summary)

footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = 1
footer.add_run("Prepared for project review and future enhancement planning")

doc.save(output_path)
print(f"Created Word file: {output_path}")
